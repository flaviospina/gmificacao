# -*- coding: utf-8 -*-
"""
estilo_gamifica.py — Kit de estilo compartilhado para as apresentações do Gamifica.
Paleta da própria plataforma (roxo #7c6ef0). Reutilizado pelos dois decks.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE

# ── Paleta ────────────────────────────────────────────────────
ROXO      = RGBColor(0x7C, 0x6E, 0xF0)
ROXO2     = RGBColor(0xA8, 0x55, 0xF7)
ESCURO    = RGBColor(0x1A, 0x1A, 0x2E)
ESCURO2   = RGBColor(0x16, 0x21, 0x3E)
CLARO     = RGBColor(0xF5, 0xF3, 0xFF)
BRANCO    = RGBColor(0xFF, 0xFF, 0xFF)
CINZA     = RGBColor(0x88, 0x88, 0x88)
CINZA_CL  = RGBColor(0xE8, 0xE8, 0xF0)
VERDE     = RGBColor(0x16, 0xA3, 0x4A)
AMBAR     = RGBColor(0xD9, 0x77, 0x06)
VERMELHO  = RGBColor(0xDC, 0x26, 0x26)
LILAS     = RGBColor(0xA7, 0x8B, 0xFA)

# Slide 16:9
LARG = Inches(13.333)
ALT  = Inches(7.5)

FONTE = "Calibri"


def nova_apresentacao():
    prs = Presentation()
    prs.slide_width = LARG
    prs.slide_height = ALT
    return prs


def slide_vazio(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fundo(slide, cor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = cor


def _sem_linha(shape):
    shape.line.fill.background()


def retangulo(slide, x, y, w, h, cor, arred=False, linha=None, linha_cor=None):
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if arred else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(forma, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = cor
    if linha and linha_cor:
        shp.line.color.rgb = linha_cor
        shp.line.width = linha
    else:
        _sem_linha(shp)
    shp.shadow.inherit = False
    return shp


def faixa_gradiente(slide, x, y, w, h, cor1, cor2, arred=True):
    """Retângulo com preenchimento gradiente diagonal."""
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if arred else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(forma, x, y, w, h)
    _sem_linha(shp)
    shp.shadow.inherit = False
    fill = shp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = cor1
    stops[1].position = 1.0
    stops[1].color.rgb = cor2
    try:
        fill.gradient_angle = 45.0
    except Exception:
        pass
    return shp


def texto(slide, x, y, w, h, txt, tam=18, cor=ESCURO, negrito=False,
          alinha=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fonte=FONTE,
          espaco=1.0, italico=False):
    cx = slide.shapes.add_textbox(x, y, w, h)
    tf = cx.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.NONE
    linhas = txt.split("\n")
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.line_spacing = espaco
        run = p.add_run()
        run.text = linha
        f = run.font
        f.size = Pt(tam)
        f.bold = negrito
        f.italic = italico
        f.name = fonte
        f.color.rgb = cor
    return cx


def bullets(slide, x, y, w, h, itens, tam=16, cor=ESCURO, marcador="•",
            cor_marcador=ROXO, espaco=1.15, gap_depois=8):
    cx = slide.shapes.add_textbox(x, y, w, h)
    tf = cx.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, item in enumerate(itens):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = espaco
        p.space_after = Pt(gap_depois)
        r1 = p.add_run()
        r1.text = f"{marcador}  "
        r1.font.size = Pt(tam)
        r1.font.bold = True
        r1.font.name = FONTE
        r1.font.color.rgb = cor_marcador
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(tam)
        r2.font.name = FONTE
        r2.font.color.rgb = cor
    return cx


def chip(slide, x, y, w, h, txt, cor_fundo, cor_txt, tam=12):
    shp = retangulo(slide, x, y, w, h, cor_fundo, arred=True)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = txt
    run.font.size = Pt(tam)
    run.font.bold = True
    run.font.name = FONTE
    run.font.color.rgb = cor_txt
    return shp


def circulo_emoji(slide, x, y, d, emoji, cor_fundo, tam=32):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shp.fill.solid()
    shp.fill.fore_color.rgb = cor_fundo
    _sem_linha(shp)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = emoji
    run.font.size = Pt(tam)
    run.font.name = "Segoe UI Emoji"
    return shp


def rodape(slide, num, total, escuro_fundo=False):
    cor = LILAS if escuro_fundo else CINZA
    texto(slide, Inches(0.5), Inches(7.02), Inches(6), Inches(0.4),
          "Gamifica · gamificação escolar sem punição", tam=10, cor=cor)
    texto(slide, Inches(11.3), Inches(7.02), Inches(1.6), Inches(0.4),
          f"{num} / {total}", tam=10, cor=cor, alinha=PP_ALIGN.RIGHT)


def logo(slide, x, y, escuro_fundo=False, tam=18):
    circulo_emoji(slide, x, y, Inches(0.5), "🎮", ROXO, tam=tam)
    texto(slide, x + Inches(0.58), y - Inches(0.02), Inches(3), Inches(0.55),
          "Gamifica", tam=20, cor=(BRANCO if escuro_fundo else ESCURO),
          negrito=True, anchor=MSO_ANCHOR.MIDDLE)
